from __future__ import annotations

from pathlib import Path
from typing import Dict, Iterable, List, Optional

from openpyxl import load_workbook

from gvai.postlabor.workers.occupation_market import (
    OccupationMarketRecord,
)
from gvai.postlabor.workers.soc_crosswalk import (
    SOCCrosswalk,
    SOCCrosswalkRecord,
)


def _clean(value):
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()

        if value in {"", "—", "-", "N/A", "na"}:
            return None

    return value


def _float(value) -> Optional[float]:
    value = _clean(value)

    if value is None:
        return None

    if isinstance(value, str):
        value = (
            value.replace(",", "")
            .replace("$", "")
            .replace("%", "")
            .strip()
        )

    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _normalize_header(value) -> str:
    if value is None:
        return ""

    return " ".join(
        str(value)
        .replace("\n", " ")
        .replace("\r", " ")
        .split()
    ).strip()


def _find_header_row(ws, required_terms: Iterable[str]) -> int:
    required = [
        term.lower()
        for term in required_terms
    ]

    for row_number in range(
        1,
        min(ws.max_row, 25) + 1,
    ):
        values = [
            _normalize_header(cell.value).lower()
            for cell in ws[row_number]
        ]

        joined = " | ".join(values)

        if all(term in joined for term in required):
            return row_number

    raise ValueError(
        "Could not detect BLS header row."
    )


def _header_map(ws, header_row: int) -> Dict[str, int]:
    result: Dict[str, int] = {}

    for cell in ws[header_row]:
        name = _normalize_header(cell.value)

        if name:
            result[name.lower()] = cell.column

    return result


def _column(
    headers: Dict[str, int],
    *contains: str,
) -> Optional[int]:
    terms = [
        item.lower()
        for item in contains
    ]

    for header, column in headers.items():
        if all(term in header for term in terms):
            return column

    return None


def load_occupational_projections(
    path: str | Path,
    *,
    sheet_name: Optional[str] = None,
    line_items_only: bool = True,
) -> List[OccupationMarketRecord]:
    """
    Load BLS Table 1.2-style occupational projections XLSX.

    Designed around the 2025-35 occupational projections schema but
    intentionally locates columns by header text rather than fixed positions.
    """

    workbook = load_workbook(
        filename=Path(path),
        read_only=True,
        data_only=True,
    )

    ws = (
        workbook[sheet_name]
        if sheet_name
        else workbook[workbook.sheetnames[0]]
    )

    header_row = _find_header_row(
        ws,
        required_terms=[
            "national employment matrix title",
            "national employment matrix code",
        ],
    )

    headers = _header_map(
        ws,
        header_row,
    )

    title_col = _column(
        headers,
        "national employment matrix title",
    )
    code_col = _column(
        headers,
        "national employment matrix code",
    )
    type_col = _column(
        headers,
        "occupation type",
    )

    emp_base_col = _column(
        headers,
        "employment",
        "2025",
    )
    emp_projected_col = _column(
        headers,
        "employment",
        "2035",
    )
    change_pct_col = _column(
        headers,
        "employment change",
        "percent",
    )
    openings_col = _column(
        headers,
        "occupational openings",
    )
    wage_col = _column(
        headers,
        "median annual wage",
    )
    education_col = _column(
        headers,
        "typical education needed",
    )
    experience_col = _column(
        headers,
        "work experience",
    )
    training_col = _column(
        headers,
        "on-the-job training",
    )

    required_columns = {
        "title": title_col,
        "code": code_col,
    }

    missing_required = [
        name
        for name, column in required_columns.items()
        if column is None
    ]

    if missing_required:
        raise ValueError(
            "Missing required BLS columns: "
            + ", ".join(missing_required)
        )

    records: List[OccupationMarketRecord] = []

    for row_number in range(
        header_row + 1,
        ws.max_row + 1,
    ):
        def val(column):
            if column is None:
                return None
            return ws.cell(
                row=row_number,
                column=column,
            ).value

        title = _clean(val(title_col))
        code = _clean(val(code_col))
        occupation_type = (
            str(_clean(val(type_col)) or "")
            .strip()
        )

        if not title or not code:
            continue

        if (
            line_items_only
            and occupation_type
            and occupation_type.lower() != "line item"
        ):
            continue

        records.append(
            OccupationMarketRecord(
                soc_code=str(code),
                title=str(title),
                employment_base_thousands=_float(
                    val(emp_base_col)
                ),
                employment_projected_thousands=_float(
                    val(emp_projected_col)
                ),
                employment_change_percent=_float(
                    val(change_pct_col)
                ),
                annual_openings_thousands=_float(
                    val(openings_col)
                ),
                median_annual_wage=_float(
                    val(wage_col)
                ),
                education=str(
                    _clean(val(education_col)) or ""
                ),
                related_experience=str(
                    _clean(val(experience_col)) or ""
                ),
                on_the_job_training=str(
                    _clean(val(training_col)) or ""
                ),
                occupation_type=occupation_type,
            )
        )

    return records


def index_by_soc(
    records: Iterable[OccupationMarketRecord],
) -> Dict[str, OccupationMarketRecord]:
    return {
        record.soc_code: record
        for record in records
    }


def load_onet_bls_crosswalk(
    path: str | Path,
    *,
    sheet_name: Optional[str] = None,
) -> SOCCrosswalk:
    """
    Load BLS O*NET-SOC -> Occupational Outlook/National Employment Matrix
    crosswalk from XLSX.

    Column detection is intentionally loose because BLS may adjust captions.
    """

    workbook = load_workbook(
        filename=Path(path),
        read_only=True,
        data_only=True,
    )

    ws = (
        workbook[sheet_name]
        if sheet_name
        else workbook[workbook.sheetnames[0]]
    )

    header_row = _find_header_row(
        ws,
        required_terms=["o*net"],
    )

    headers = _header_map(
        ws,
        header_row,
    )

    onet_code_col = (
        _column(headers, "o*net", "code")
        or _column(headers, "o*net-soc")
    )
    onet_title_col = _column(
        headers,
        "o*net",
        "title",
    )

    bls_code_col = (
        _column(headers, "matrix", "code")
        or _column(headers, "ooh", "code")
        or _column(headers, "soc", "code")
    )

    bls_title_col = (
        _column(headers, "matrix", "title")
        or _column(headers, "ooh", "title")
    )

    if onet_code_col is None:
        raise ValueError(
            "Could not identify O*NET-SOC code column."
        )

    if bls_code_col is None:
        raise ValueError(
            "Could not identify BLS SOC/NEM code column."
        )

    records: List[SOCCrosswalkRecord] = []

    for row_number in range(
        header_row + 1,
        ws.max_row + 1,
    ):
        def val(column):
            if column is None:
                return None
            return ws.cell(
                row=row_number,
                column=column,
            ).value

        onet_code = _clean(
            val(onet_code_col)
        )
        bls_code = _clean(
            val(bls_code_col)
        )

        if not onet_code or not bls_code:
            continue

        records.append(
            SOCCrosswalkRecord(
                onet_soc_code=str(onet_code),
                bls_soc_code=str(bls_code),
                onet_title=str(
                    _clean(val(onet_title_col)) or ""
                ),
                bls_title=str(
                    _clean(val(bls_title_col)) or ""
                ),
            )
        )

    return SOCCrosswalk(records)
