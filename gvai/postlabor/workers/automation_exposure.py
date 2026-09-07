from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from openpyxl import load_workbook


AI_CATEGORY_SCORE = {
    "low": 20.0,
    "moderate": 45.0,
    "high": 70.0,
    "very high": 90.0,
}


@dataclass(frozen=True)
class AIExposureRecord:
    soc_code: str
    title: str
    category: str
    theoretical_percentile: Optional[float] = None
    observed_percentile: Optional[float] = None

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)


@dataclass(frozen=True)
class AutomationAssessment:
    soc_code: str
    title: str
    ai_exposure_score: float
    displacement_pressure: float
    augmentation_potential: float
    employment_change_percent: Optional[float]
    confidence: float
    explanation: List[str]

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)


def _clean(value):
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()

        if value in {"", "-", "—", "N/A"}:
            return None

    return value


def _float(value) -> Optional[float]:
    value = _clean(value)

    if value is None:
        return None

    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _normalize_header(value) -> str:
    if value is None:
        return ""

    return " ".join(
        str(value)
        .replace("\n", " ")
        .replace("\r", " ")
        .split()
    ).strip().lower()


def load_ai_exposure(
    path: str | Path,
) -> List[AIExposureRecord]:
    workbook = load_workbook(
        filename=Path(path),
        read_only=False,
        data_only=True,
    )

    ws = workbook[workbook.sheetnames[0]]

    header_row = None

    for row_number in range(1, min(ws.max_row, 30) + 1):
        headers = [
            _normalize_header(cell.value)
            for cell in ws[row_number]
        ]

        joined = " | ".join(headers)

        if "occupation" in joined and "exposure" in joined:
            header_row = row_number
            break

    if header_row is None:
        raise ValueError("Could not detect BLS AI exposure header row.")

    headers: Dict[str, int] = {}

    for cell in ws[header_row]:
        name = _normalize_header(cell.value)

        if name:
            headers[name] = cell.column

    def find_col(*terms: str) -> Optional[int]:
        wanted = [term.lower() for term in terms]

        for header, column in headers.items():
            if all(term in header for term in wanted):
                return column

        return None

    code_col = (
        find_col("code")
        or find_col("soc")
    )

    title_col = (
        find_col("title")
        or find_col("occupation")
    )

    category_col = (
        find_col("relative", "ai", "exposure")
        or find_col("exposure", "category")
    )

    # The current BLS workbook publishes the four-category
    # Relative AI Exposure classification directly. Percentile fields
    # are optional and may not exist in the distributed XLSX.
    theoretical_col = (
        find_col("theoretical")
        or find_col("theoretical", "percentile")
    )

    observed_col = (
        find_col("observed")
        or find_col("current", "evidence")
    )

    if code_col is None:
        raise ValueError("Could not locate occupation code column.")

    if category_col is None:
        raise ValueError("Could not locate AI exposure category column.")

    records: List[AIExposureRecord] = []

    for row_number in range(header_row + 1, ws.max_row + 1):

        def val(column):
            if column is None:
                return None

            return ws.cell(
                row=row_number,
                column=column,
            ).value

        code = _clean(val(code_col))
        category = _clean(val(category_col))

        if not code or not category:
            continue

        records.append(
            AIExposureRecord(
                soc_code=str(code).strip(),
                title=str(_clean(val(title_col)) or ""),
                category=str(category).strip(),
                theoretical_percentile=_float(
                    val(theoretical_col)
                ),
                observed_percentile=_float(
                    val(observed_col)
                ),
            )
        )

    return records


def index_ai_exposure(
    records: Iterable[AIExposureRecord],
) -> Dict[str, AIExposureRecord]:
    return {
        record.soc_code: record
        for record in records
    }


def ai_exposure_score(
    record: AIExposureRecord,
) -> float:
    category = record.category.strip().lower()

    base = AI_CATEGORY_SCORE.get(
        category,
        50.0,
    )

    # If percentile information is available, blend it in.
    percentiles = [
        value
        for value in (
            record.theoretical_percentile,
            record.observed_percentile,
        )
        if value is not None
    ]

    if not percentiles:
        return round(base, 2)

    average_percentile = (
        sum(percentiles)
        / len(percentiles)
    )

    if average_percentile <= 1.0:
        average_percentile *= 100.0

    return round(
        base * 0.50
        + average_percentile * 0.50,
        2,
    )


def assess_automation(
    *,
    exposure: AIExposureRecord,
    employment_change_percent: Optional[float],
    physical_task_resilience: float = 50.0,
    augmentation_potential: float = 50.0,
    confidence: float = 0.75,
) -> AutomationAssessment:
    """
    Estimate displacement pressure separately from AI exposure.

    v0.1 combines:
    - BLS relative AI exposure
    - employment trajectory
    - physical/task resilience
    - augmentation potential

    High AI exposure alone does NOT imply displacement.
    """

    exposure_score = ai_exposure_score(exposure)

    if employment_change_percent is None:
        decline_pressure = 50.0
    else:
        # +20% growth -> 0 pressure
        # -20% decline -> 100 pressure
        decline_pressure = (
            (20.0 - employment_change_percent)
            / 40.0
        ) * 100.0

        decline_pressure = max(
            0.0,
            min(100.0, decline_pressure),
        )

    physical_vulnerability = (
        100.0 - physical_task_resilience
    )

    augmentation_resistance = (
        100.0 - augmentation_potential
    )

    displacement = (
        exposure_score * 0.40
        + decline_pressure * 0.25
        + physical_vulnerability * 0.20
        + augmentation_resistance * 0.15
    )

    explanation = [
        f"BLS relative AI exposure: {exposure.category}.",
        f"AI exposure score: {exposure_score:.1f}/100.",
        (
            "Employment trend contributes separately from AI exposure: "
            f"{employment_change_percent}%."
            if employment_change_percent is not None
            else "Employment trend unavailable."
        ),
        (
            "Physical/task resilience: "
            f"{physical_task_resilience:.1f}/100."
        ),
        (
            "Augmentation potential: "
            f"{augmentation_potential:.1f}/100."
        ),
        (
            "AI exposure is treated as exposure, not as a direct "
            "probability of job loss."
        ),
    ]

    return AutomationAssessment(
        soc_code=exposure.soc_code,
        title=exposure.title,
        ai_exposure_score=round(exposure_score, 2),
        displacement_pressure=round(displacement, 2),
        augmentation_potential=round(
            augmentation_potential,
            2,
        ),
        employment_change_percent=(
            round(employment_change_percent, 2)
            if employment_change_percent is not None
            else None
        ),
        confidence=round(confidence, 4),
        explanation=explanation,
    )
